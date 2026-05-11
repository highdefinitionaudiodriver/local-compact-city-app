"""Build docs/設計書.xlsx from the same source content as 設計書.md.

This is a one-shot generator. It does not parse the markdown; instead the
content is structured directly so each sheet is hand-tuned for legibility.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── Styles ─────────────────────────────────────────────────────────────
FONT_JP = "Yu Gothic"
MONO = "Consolas"

TITLE      = Font(name=FONT_JP, size=22, bold=True, color="1E3A8A")
SUBTITLE   = Font(name=FONT_JP, size=11, italic=True, color="64748B")
SECTION    = Font(name=FONT_JP, size=14, bold=True, color="FFFFFF")
SECTION_F  = PatternFill("solid", start_color="1E3A8A")
SUBSECTION = Font(name=FONT_JP, size=12, bold=True, color="1E3A8A")
SUBSEC_F   = PatternFill("solid", start_color="DBEAFE")
TH         = Font(name=FONT_JP, size=10, bold=True, color="FFFFFF")
TH_F       = PatternFill("solid", start_color="475569")
TD         = Font(name=FONT_JP, size=10, color="0F172A")
TD_ALT_F   = PatternFill("solid", start_color="F8FAFC")
BODY       = Font(name=FONT_JP, size=10, color="0F172A")
NOTE       = Font(name=FONT_JP, size=9,  italic=True, color="64748B")
MONO_F     = Font(name=MONO,   size=9,  color="0F172A")
LINK       = Font(name=FONT_JP, size=10, color="1D4ED8", underline="single")

THIN   = Side(style="thin",   color="CBD5E1")
MEDIUM = Side(style="medium", color="334155")
B_TD = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
B_TH = Border(top=MEDIUM, bottom=MEDIUM, left=THIN, right=THIN)

WRAP_TOP   = Alignment(wrap_text=True, vertical="top",    horizontal="left")
WRAP_CTR   = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_LEFT  = Alignment(wrap_text=True, vertical="center", horizontal="left")

wb = Workbook()
wb.remove(wb.active)


def new_sheet(title, tab="1E3A8A"):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    return ws


def widths(ws, *ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def hdr(ws, row, text, span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill = SECTION, SECTION_F
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 28
    return row + 1


def sub(ws, row, text, span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill = SUBSECTION, SUBSEC_F
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22
    return row + 1


def para(ws, row, text, span, font=BODY, height=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font
    c.alignment = WRAP_TOP
    if height is not None:
        ws.row_dimensions[row].height = height
    return row + 1


def table(ws, start, headers, rows, alt=True, row_height=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.font, c.fill = TH, TH_F
        c.alignment = WRAP_CTR
        c.border = B_TH
    ws.row_dimensions[start].height = 22
    for i, r in enumerate(rows):
        rn = start + 1 + i
        for j, v in enumerate(r, 1):
            c = ws.cell(row=rn, column=j, value=v)
            c.font = TD
            c.alignment = WRAP_TOP
            c.border = B_TD
            if alt and i % 2 == 1:
                c.fill = TD_ALT_F
        if row_height is not None:
            ws.row_dimensions[rn].height = row_height
    return start + 1 + len(rows)


def gap(row, n=1):
    return row + n


# ════════════════════════════════════════════════════════════════════
# Sheet 1: 表紙
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("表紙", "1E3A8A")
widths(ws, 6, 26, 56, 6)

ws.merge_cells("B2:C3")
c = ws.cell(row=2, column=2, value="ローカル・コンパクトシティ")
c.font = TITLE
c.alignment = Alignment(vertical="bottom", horizontal="left")
ws.row_dimensions[2].height = 30
ws.row_dimensions[3].height = 12

ws.merge_cells("B4:C4")
c = ws.cell(row=4, column=2, value="住民提案・デジタル署名プラットフォーム 設計書")
c.font = Font(name=FONT_JP, size=15, bold=True, color="0F172A")
ws.row_dimensions[4].height = 26

ws.merge_cells("B5:C5")
c = ws.cell(row=5, column=2,
    value="マイナンバーカード（JPKI）認証で住民であることを証明し、改ざん検知可能な公文書として行政に届ける")
c.font = SUBTITLE

r = 8
r = sub(ws, r, "文書情報", 3)
r = table(ws, r,
    ["項目", "内容", ""],
    [
        ["文書名", "ローカル・コンパクトシティ 住民提案・デジタル署名プラットフォーム 設計書", ""],
        ["バージョン", "1.0", ""],
        ["作成日", "2026-05-10", ""],
        ["対象システム", "プロトタイプ / コンセプト実証 (PoC)", ""],
        ["想定読者", "自治体 IT 担当・評価ベンダー・開発引継ぎ要員・プロジェクトマネージャ", ""],
        ["関連文書", "README.md (デモ脚本) / prisma/schema.prisma (DB正本)", ""],
    ],
    row_height=22,
)
# Hide trailing column header
ws.cell(row=r-7, column=3).value = None
ws.cell(row=r-7, column=3).fill = PatternFill()
ws.cell(row=r-7, column=3).border = Border()

r = gap(r, 2)
r = sub(ws, r, "本書の構成", 3)
r = table(ws, r,
    ["章", "節タイトル", "概要"],
    [
        ["1", "はじめに", "文書目的・読者・関連文書"],
        ["2", "システム概要", "解決課題・スコープ・ユースケース"],
        ["3", "アーキテクチャ", "全体構成・技術選定・レイヤ分離"],
        ["4", "データモデル", "ER 概要・テーブル詳細"],
        ["5", "機能仕様", "13 機能の入出力・制約"],
        ["6", "画面設計", "11 画面・遷移・主要画面構成"],
        ["7", "API / Server Action 仕様", "Server Actions・Route Handlers"],
        ["8", "認証・認可設計", "JPKI モック・xID 連携・ロール"],
        ["9", "マルチテナント設計", "テナント分離の各層"],
        ["10", "セキュリティ設計", "個人情報最小化・改ざん検知"],
        ["11", "監査・証跡設計", "三層監査・append-only"],
        ["12", "非機能要件", "性能・可用性・拡張性"],
        ["13", "デプロイ・運用設計", "環境構成・環境変数・バックアップ"],
        ["14", "既知の制約と本番化ロードマップ", "制約一覧・5 フェーズ計画"],
        ["付録 A", "用語集", "JPKI/HMAC/OIDC 等"],
        ["付録 B", "自治体マスタ", "現プロト登録自治体"],
        ["付録 C", "改訂履歴", "版管理"],
    ],
    row_height=20,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 2: 目次
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("目次", "1E3A8A")
widths(ws, 6, 34, 66)

r = 2
r = hdr(ws, r, "目次", 3)
r = gap(r)
r = para(ws, r, "各章は同名のシートに対応します。シートタブから直接ジャンプできます。", 3, NOTE)
r = gap(r)
r = table(ws, r,
    ["章", "シート名", "扱う主要トピック"],
    [
        ["1",       "1.はじめに",          "目的・読者・関連文書"],
        ["2",       "2.システム概要",       "課題・スコープ・ユースケース"],
        ["3",       "3.アーキテクチャ",     "構成図・技術選定理由"],
        ["4",       "4.データモデル",       "ER・テーブル詳細"],
        ["5",       "5.機能仕様",          "13 機能の入出力・制約"],
        ["6",       "6.画面設計",          "画面一覧・遷移・構成"],
        ["7",       "7.API仕様",           "Server Action / Route Handler"],
        ["8",       "8.認証認可",          "JPKI / xID / ロール"],
        ["9",       "9.マルチテナント",     "テナント識別・分離・越境"],
        ["10",      "10.セキュリティ",     "個人情報最小化・HMAC"],
        ["11",      "11.監査証跡",         "三層監査・append-only"],
        ["12",      "12.非機能要件",       "性能・可用性・スケール"],
        ["13",      "13.デプロイ運用",     "環境・変数・バックアップ"],
        ["14",      "14.制約とRM",         "制約・本番化ロードマップ"],
        ["付録 A",  "付録",                "用語集・自治体マスタ・改訂履歴"],
    ],
    row_height=20,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 3: 1.はじめに
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("1.はじめに", "3B82F6")
widths(ws, 24, 76)

r = 2
r = hdr(ws, r, "1. はじめに", 2)
r = gap(r)
r = sub(ws, r, "1.1 文書の目的", 2)
r = para(ws, r,
    "本書は、地域住民が自治体に対してマイナンバーカード（JPKI）認証付きの電子署名で請願・提案を行うプラットフォームの設計を、"
    "開発・運用・評価の各関係者が共通理解を持てる粒度で記述する。",
    2, BODY, 48)
r = gap(r)
r = sub(ws, r, "1.2 対象読者", 2)
r = table(ws, r,
    ["読者", "主に参照する章"],
    [
        ["自治体 IT 担当・調達担当", "1, 2, 8, 10, 11, 13, 14"],
        ["評価ベンダー・監査人",     "4, 7, 8, 10, 11"],
        ["開発担当・引継ぎ要員",     "3, 4, 5, 6, 7, 9, 13"],
        ["プロジェクトマネージャ",   "1, 2, 12, 14"],
    ],
    row_height=22,
)
r = gap(r)
r = sub(ws, r, "1.3 関連文書", 2)
r = table(ws, r,
    ["文書", "役割"],
    [
        ["README.md",                       "デモ実演台本・セットアップ手順"],
        ["prisma/schema.prisma",            "DB スキーマ正本"],
        ["prisma/migrations/",              "DDL 履歴（不可逆）"],
        ["docs/設計書.md",                  "本設計書（Markdown 版）"],
    ],
    row_height=22,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 4: 2.システム概要
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("2.システム概要", "3B82F6")
widths(ws, 38, 60)

r = 2
r = hdr(ws, r, "2. システム概要", 2)
r = gap(r)
r = sub(ws, r, "2.1 解決する課題", 2)
r = table(ws, r,
    ["行政・社会課題", "本システムの解決アプローチ"],
    [
        ["紙ベースの請願は集約・保管コストが高い", "デジタル投稿 + 自動集計 + 公文書品質 PDF 出力"],
        ["ネット署名は本人性が証明できない",       "JPKI による居住確認済み電子署名"],
        ["配布された電子文書の改ざんを検知できない", "HMAC-SHA256 を文書に埋め込み、サーバ側で再検証可能"],
        ["行政が市民の関心を肌感で把握できない",   "統計ダッシュボード（時系列・属性別）"],
        ["マイナンバーの取扱い不安",                "マイナンバー本体は一切取得しない設計"],
    ],
    row_height=36,
)
r = gap(r)

r = sub(ws, r, "2.2 スコープ", 2)
r = table(ws, r,
    ["区分", "内容"],
    [
        ["スコープ内",
         "・住民による「アイデア (IDEA)」「請願 (PETITION)」の投稿\n"
         "・マイナンバーカード認証（モック）\n"
         "・賛否表明・電子署名・コメント\n"
         "・自治体職員向け管理画面（ステータス管理）\n"
         "・CSV / PDF エクスポート（改ざん検知対応）\n"
         "・統計ダッシュボード\n"
         "・原本性検証ページ\n"
         "・マルチテナント（複数自治体共存）\n"
         "・監査ログ・発行台帳・検証履歴"],
        ["スコープ外",
         "・実機 JPKI 連携（接続スタブのみ用意）\n"
         "・メール / プッシュ通知\n"
         "・多言語化（i18n）\n"
         "・モバイルアプリ\n"
         "・公開鍵署名（ECDSA / Ed25519）への移行\n"
         "・WORM ストレージ統合"],
    ],
    row_height=160,
)
r = gap(r)

r = sub(ws, r, "2.3 主要ユースケース", 2)
r = table(ws, r,
    ["ID / アクター", "ユースケース"],
    [
        ["UC-01 住民",    "マイナンバーカードで認証する"],
        ["UC-02 住民",    "自治体に提案・請願を投稿する"],
        ["UC-03 住民",    "提案に賛同・反対する (IDEA)"],
        ["UC-04 住民",    "請願に電子署名する (PETITION)"],
        ["UC-05 住民",    "提案にコメントを投稿する"],
        ["UC-06 自治体職員", "自治体内の提案を一覧・管理する"],
        ["UC-07 自治体職員", "提案のステータスを変更する"],
        ["UC-08 自治体職員", "CSV / PDF 報告書を発行する"],
        ["UC-09 自治体職員", "統計ダッシュボードで活動量を把握する"],
        ["UC-10 他部署職員", "受け取った CSV / PDF の原本性を検証する"],
    ],
    row_height=20,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 5: 3.アーキテクチャ
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("3.アーキテクチャ", "8B5CF6")
widths(ws, 30, 28, 50)

r = 2
r = hdr(ws, r, "3. アーキテクチャ", 3)
r = gap(r)

r = sub(ws, r, "3.1 全体構成（テキスト表現）", 3)
r = para(ws, r,
    "クライアント (Browser) → Edge Middleware → Node Runtime (Server Actions / Route Handlers) → Data Layer (Prisma + SQLite)\n"
    "Recharts などのクライアント可視化は Server Component で集計した props を受け取る。\n"
    "PDF 生成は Node Runtime で react-pdf renderer + バンドル済 Noto Sans JP を用いる。\n"
    "Edge レイヤには node:crypto 等の Node 専用 API は持ち込まない（auth.config.ts のみ）。",
    3, BODY, 100)
r = gap(r)

r = sub(ws, r, "3.2 技術選定", 3)
r = table(ws, r,
    ["レイヤ", "採用技術", "選定理由"],
    [
        ["Frontend",       "Next.js 16 (App Router, Turbopack)", "Server Components で初期表示の軽量化、Server Actions で API 層を薄く"],
        ["UI ライブラリ",   "shadcn/ui + Tailwind v4",           "コード化されたコンポーネントでカスタマイズ容易"],
        ["言語",            "TypeScript 5",                       "型による契約、リファクタリング安全性"],
        ["Auth",            "NextAuth.js v5 (Auth.js)",           "OIDC 対応、JWT/DB セッション切替自由"],
        ["ORM",             "Prisma 6",                           "型安全、マイグレーション管理、トランザクション API"],
        ["DB",              "SQLite → PostgreSQL",                 "プロト容易、本番は接続文字列のみ変更"],
        ["PDF",             "@react-pdf/renderer v4",             "React コンポーネントで PDF を組める"],
        ["グラフ",           "Recharts v3",                        "宣言的、SVG ベース、Server props と相性◎"],
        ["PDF 解析",         "pdf-lib",                            "metadata 読取が軽量、純 JS"],
    ],
    row_height=28,
)
r = gap(r)

r = sub(ws, r, "3.3 レイヤ分離方針", 3)
r = table(ws, r,
    ["レイヤ", "責務", "代表ファイル"],
    [
        ["Edge",        "URL ベースの認可ガードのみ。crypto は持ち込まない", "src/proxy.ts, src/auth.config.ts"],
        ["Node Server", "DB アクセス、PDF 生成、HMAC 計算、JPKI 認証",       "src/auth.ts, Route Handlers"],
        ["Client",      "ユーザ操作、ローカル状態、可視化",                   "login-form.tsx, charts.tsx"],
        ["Pure Data",   "自治体マスタ、enum 定数",                           "src/lib/constants.ts"],
    ],
    row_height=32,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 6: 4.データモデル
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("4.データモデル", "10B981")
widths(ws, 22, 18, 18, 60)

r = 2
r = hdr(ws, r, "4. データモデル", 4)
r = gap(r)

r = sub(ws, r, "4.1 ER 関係", 4)
r = para(ws, r,
    "User 1 — N Proposal  (User.id → Proposal.author_id)\n"
    "User 1 — N Vote      (User.id → Vote.user_id)\n"
    "User 1 — N Comment   (User.id → Comment.user_id)\n"
    "Proposal 1 — N Vote     (Proposal.id → Vote.proposal_id, ON DELETE CASCADE)\n"
    "Proposal 1 — N Comment  (Proposal.id → Comment.proposal_id, ON DELETE CASCADE)\n"
    "Proposal 1 — N AuditLog (Proposal.id → AuditLog.proposal_id)\n"
    "IssuanceRecord, VerificationLog は append-only かつ独立（FK 参照は緩い）",
    4, MONO_F, 120)
r = gap(r)

# 4.2 各テーブル詳細
def write_columns(ws, r, title, rows):
    r = sub(ws, r, title, 4)
    r = table(ws, r,
        ["カラム", "型", "制約", "説明"],
        rows,
        row_height=28,
    )
    return r

r = write_columns(ws, r, "User", [
    ["id",             "String",   "PK",          "sha256(certificate_serial) — 不可逆ハッシュ"],
    ["resident_code",  "String",   "NOT NULL",    "6 桁自治体コード（JIS X 0402 ベース）"],
    ["display_name",   "String?",  "NULL 可",     "JPKI から取得した氏名（任意）"],
    ["is_verified",    "Boolean",  "DEFAULT true", "JPKI 認証経由なら常に true"],
    ["created_at",     "DateTime", "DEFAULT now()", ""],
])
r = gap(r)

r = write_columns(ws, r, "Proposal", [
    ["id",                "cuid",     "PK", ""],
    ["title, content",    "String",   "NOT NULL", "提案内容"],
    ["type",              "String",   "",          "IDEA / PETITION"],
    ["status",            "String",   "",          "OPEN / CLOSED / SUBMITTED"],
    ["target_signatures", "Int",      "DEFAULT 100", "PETITION の目標署名数"],
    ["deadline",          "DateTime?", "",         "任意"],
    ["resident_code",     "String",   "NOT NULL",  "★テナント識別子"],
    ["author_id",         "String",   "FK→User.id", ""],
])
r = gap(r)

r = write_columns(ws, r, "Vote ( @@unique([proposal_id, user_id]) )", [
    ["id",            "cuid",     "PK", ""],
    ["proposal_id",   "String",   "FK CASCADE", ""],
    ["user_id",       "String",   "FK CASCADE", ""],
    ["vote_type",     "String",   "",        "FOR / AGAINST / SIGN"],
    ["resident_code", "String",   "NOT NULL", "署名時点のスナップショット"],
    ["created_at",    "DateTime", "DEFAULT now()", ""],
])
r = gap(r)

r = write_columns(ws, r, "Comment", [
    ["id",          "cuid",     "PK", ""],
    ["text",        "String",   "NOT NULL", "アプリ層で 400 文字制約"],
    ["proposal_id", "String",   "FK", ""],
    ["user_id",     "String",   "FK", "UI には出さない"],
    ["created_at",  "DateTime", "DEFAULT now()", ""],
])
r = gap(r)

r = write_columns(ws, r, "AuditLog (append-only)", [
    ["id",             "cuid",     "PK", ""],
    ["proposal_id",    "String",   "FK", ""],
    ["admin_user_id",  "String",   "NOT NULL", "操作した職員のハッシュ ID"],
    ["old_status",     "String",   "NOT NULL", ""],
    ["new_status",     "String",   "NOT NULL", ""],
    ["created_at",     "DateTime", "DEFAULT now()", ""],
])
r = gap(r)

r = write_columns(ws, r, "IssuanceRecord (append-only)", [
    ["id",              "cuid",     "PK", ""],
    ["kind",            "String",   "",          "CSV / PDF"],
    ["proposal_id",     "String",   "",          ""],
    ["issuer_user_id",  "String",   "",          "発行した職員のハッシュ ID"],
    ["hmac",            "String",   "UNIQUE",    "文書の HMAC-SHA256 ダイジェスト"],
    ["payload_json",    "String",   "",          "HMAC の入力となった正規化 JSON"],
    ["issued_at",       "DateTime", "DEFAULT now()", ""],
])
r = gap(r)

r = write_columns(ws, r, "VerificationLog (append-only)", [
    ["id",                "cuid",     "PK", ""],
    ["verifier_user_id",  "String",   "NOT NULL", "検証した職員のハッシュ ID"],
    ["kind",              "String",   "",          "CSV / PDF / UNKNOWN"],
    ["result",            "String",   "",          "OK / TAMPERED / UNKNOWN"],
    ["hmac",              "String?",  "NULL 可",   "ファイルから抽出した HMAC"],
    ["file_name",         "String?",  "",          ""],
    ["file_size",         "Int?",     "",          ""],
    ["created_at",        "DateTime", "DEFAULT now()", ""],
])

# ════════════════════════════════════════════════════════════════════
# Sheet 7: 5.機能仕様
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("5.機能仕様", "F59E0B")
widths(ws, 8, 26, 16, 30, 30)

r = 2
r = hdr(ws, r, "5. 機能仕様", 5)
r = gap(r)
r = sub(ws, r, "5.1 機能一覧", 5)
r = table(ws, r,
    ["ID", "機能", "アクター", "入力", "主な制約"],
    [
        ["F-01", "JPKI 認証ログイン",         "住民/職員", "ペルソナ選択 / xID OIDC", "マイナンバー本体は取得しない"],
        ["F-02", "提案投稿",                  "住民",     "タイトル・本文・種別・目標",  "種別ごとに項目バリデーション"],
        ["F-03", "賛否表明 (IDEA)",           "住民",     "提案 ID + FOR/AGAINST",       "自治体一致、二重防止"],
        ["F-04", "電子署名 (PETITION)",       "住民",     "提案 ID + 確認チェック",      "自治体一致、同意宣言必須"],
        ["F-05", "コメント投稿",               "住民",     "提案 ID + 本文(400字)",        "trim 後 1〜400 字"],
        ["F-06", "提案管理ダッシュボード",     "職員",     "—",                            "自治体スコープで自動絞り込み"],
        ["F-07", "ステータス変更",             "職員",     "提案 ID + 新ステータス",       "トランザクション + 監査ログ"],
        ["F-08", "CSV エクスポート",          "職員",     "提案 ID",                      "PETITION のみ、HMAC 付与"],
        ["F-09", "PDF レポート",              "職員",     "提案 ID",                      "PETITION のみ、metadata + 印字"],
        ["F-10", "統計ダッシュボード",         "職員",     "—",                            "テナントスコープ"],
        ["F-11", "原本性検証",                "職員",     "アップロード ファイル",        "3 段階照合"],
        ["F-12", "行政レポート(公開)",         "全員",     "—",                            "公開請願の集計表示"],
        ["F-13", "xID 連携(スタブ)",           "住民",     "OIDC コールバック",            "env で有効化"],
    ],
    row_height=28,
)
r = gap(r)

r = sub(ws, r, "5.2 主要機能の処理フロー", 5)
r = para(ws, r, "■ F-04 電子署名（PETITION）", 5, SUBSECTION, 22)
r = para(ws, r,
    "1. ユーザが該当ページで「住民として賛同します」チェックを ON\n"
    "2. 「電子署名する」ボタン押下 → Server Action castVoteAction 呼出\n"
    "3. 認証チェック（session 取得）\n"
    "4. zod でフォーム検証\n"
    "5. Proposal を取得し以下を確認:\n"
    "   ・ 存在\n"
    "   ・ status === \"OPEN\"\n"
    "   ・ 自治体一致（proposal.resident_code === session.user.resident_code）\n"
    "   ・ 種別と vote_type の整合\n"
    "   ・ 確認チェック attest === \"1\"\n"
    "   ・ 締切未経過\n"
    "6. prisma.vote.create\n"
    "   ├ 成功: revalidatePath → UI 更新\n"
    "   └ P2002: 「既に署名済み」を返却",
    5, MONO_F, 220)
r = gap(r)

r = para(ws, r, "■ F-09 PDF レポート発行", 5, SUBSECTION, 22)
r = para(ws, r,
    "1. 職員が /admin/proposals/[id] で「PDF報告書」リンク\n"
    "2. /api/admin/proposals/[id]/pdf にアクセス\n"
    "3. isAdmin(session) + 自治体一致 + PETITION 限定チェック\n"
    "4. Proposal + 関連 Vote(SIGN) を取得し集計\n"
    "5. issuedAt = new Date() を確定\n"
    "6. payloadCanonical = canonicalize(buildHmacPayload(...))\n"
    "7. hmac = hmacHex(payloadCanonical)\n"
    "8. IssuanceRecord.upsert（hmac で一意）\n"
    "9. renderToBuffer(<PetitionReport ... />)\n"
    "10. Content-Type: application/pdf でストリーム返却",
    5, MONO_F, 180)
r = gap(r)

r = para(ws, r, "■ F-11 原本性検証", 5, SUBSECTION, 22)
r = para(ws, r,
    "CSV:\n"
    " 1. 末尾 # HMAC-SHA256: <hex> 行をプレフィックス緩マッチで検出\n"
    " 2. 値が 64 桁 hex かを別途検証 → NG: TAMPERED / OK: 次へ\n"
    " 3. 行を除いた body を hmacHex() で再計算\n"
    " 4. safeEqualHex() で定数時間比較 → 不一致: TAMPERED / 一致: 次へ\n"
    " 5. IssuanceRecord をヒット試行\n"
    " 6. VerificationLog.create → OK 結果返却\n"
    "\n"
    "PDF:\n"
    " 1. pdf-lib で Keywords metadata 取得\n"
    " 2. hmac=<hex> 抽出\n"
    " 3. IssuanceRecord.findUnique({where:{hmac}}) → 不在: TAMPERED\n"
    " 4. payload_json から HMAC 再計算\n"
    " 5. safeEqualHex で比較 → 不一致: TAMPERED / 一致: OK\n"
    " 6. VerificationLog.create",
    5, MONO_F, 280)

# ════════════════════════════════════════════════════════════════════
# Sheet 8: 6.画面設計
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("6.画面設計", "F59E0B")
widths(ws, 32, 26, 18, 22)

r = 2
r = hdr(ws, r, "6. 画面設計", 4)
r = gap(r)
r = sub(ws, r, "6.1 画面一覧", 4)
r = table(ws, r,
    ["パス", "画面名", "認可", "種別"],
    [
        ["/",                              "提案一覧（ホーム）",        "公開 / 認証で絞込",  "住民向け"],
        ["/login",                         "ログイン",                  "公開",               "住民向け"],
        ["/proposals/new",                 "提案投稿",                  "認証必須",           "住民向け"],
        ["/proposals/[id]",                "提案詳細 + コメント",       "公開（投票は認証）", "住民向け"],
        ["/report",                        "行政レポート一覧",          "公開",               "公開"],
        ["/report/[id]",                   "行政レポート詳細",          "公開",               "公開"],
        ["/admin/proposals",               "管理ダッシュボード",        "職員のみ",           "行政"],
        ["/admin/proposals/[id]",          "個別管理 + 履歴",          "職員のみ",           "行政"],
        ["/admin/analytics",               "統計",                      "職員のみ",           "行政"],
        ["/admin/verify",                  "原本性検証",                "職員のみ",           "行政"],
        ["/admin/forbidden",               "非職員フォールバック",      "認証必須",           "行政"],
    ],
    row_height=22,
)
r = gap(r)

r = sub(ws, r, "6.2 画面遷移", 4)
r = para(ws, r,
    "Home (/) → Login (/login)\n"
    "Home → ProposalDetail (/proposals/:id)\n"
    "Login →(住民)→ Home\n"
    "Login →(職員)→ /admin/proposals\n"
    "Home → /proposals/new → ProposalDetail\n"
    "ProposalDetail →(職員のみ表示)→ /admin/proposals/:id\n"
    "/admin/proposals → /admin/proposals/:id / /admin/analytics / /admin/verify\n"
    "/admin/proposals/:id → CSV (Route) / PDF (Route)\n"
    "/admin/verify → (uploads) → /admin/verify",
    4, MONO_F, 180)
r = gap(r)

r = sub(ws, r, "6.3 ログイン画面の段階表示", 4)
r = para(ws, r,
    "1. ペルソナ選択（匿名ラベルのみ表示、氏名は出さない）\n"
    "2. 「カードを読み取る」ボタン → 約 1.8 秒のリーダーアニメ\n"
    "   ✓ カードを検出しました\n"
    "   ✓ PIN を検証中（モック：自動入力）...\n"
    "   ✓ 署名用電子証明書を読み取り中...\n"
    "   ✓ 認証完了\n"
    "3. 緑ボックスで「読み取った氏名」を初めて表示\n"
    "4. 「ログインしてダッシュボードへ →」を明示クリック → / へ遷移\n"
    "（自動遷移しない。プレゼン中に観客に認知を促す目的）",
    4, BODY, 200)

# ════════════════════════════════════════════════════════════════════
# Sheet 9: 7.API仕様
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("7.API仕様", "F59E0B")
widths(ws, 30, 28, 22, 32)

r = 2
r = hdr(ws, r, "7. API / Server Action 仕様", 4)
r = gap(r)
r = sub(ws, r, "7.1 Server Actions", 4)
r = table(ws, r,
    ["関数", "入力", "出力", "主な失敗"],
    [
        ["jpkiLoginAction(formData)",                "personaKey",                       "redirect",                  "persona 不在"],
        ["logoutAction()",                            "—",                                "redirect",                  "—"],
        ["createProposalAction(prev, formData)",      "title, content, type, target, deadline", "{ok?, error?, fieldErrors?}", "認証なし、バリデーションエラー"],
        ["castVoteAction(prev, formData)",            "proposalId, voteType, attest",     "{ok?, error?}",             "認証なし、自治体不一致、二重投票 (P2002)、締切経過"],
        ["postCommentAction(prev, formData)",         "proposalId, text",                 "{ok?, error?}",             "文字数超過、空文字"],
        ["updateProposalStatusAction(prev, formData)","proposalId, status",               "{ok?, error?}",             "非職員、他自治体、トランザクション失敗"],
        ["verifyExportAction(formData)",              "file",                             "VerifyResult",              "非職員、サイズ超過、形式不明"],
    ],
    row_height=42,
)
r = gap(r)

r = sub(ws, r, "7.2 Route Handlers", 4)
r = table(ws, r,
    ["メソッド + パス", "認可", "入力", "出力"],
    [
        ["GET /api/admin/proposals/[id]/export", "職員", "path id", "text/csv (HMAC 付き)"],
        ["GET /api/admin/proposals/[id]/pdf",    "職員", "path id", "application/pdf (HMAC 付き)"],
        ["GET/POST /api/auth/[...nextauth]",     "—",   "NextAuth 標準", "リダイレクト等"],
    ],
    row_height=24,
)
r = gap(r)

r = sub(ws, r, "7.3 castVoteAction のレスポンス例", 4)
r = para(ws, r,
    "成功                : { \"ok\": true }\n"
    "自治体不一致       : { \"error\": \"他自治体の提案には署名・投票できません。お住まいの自治体の提案のみご参加いただけます。\" }\n"
    "二重投票           : { \"error\": \"この提案には既に投票/署名済みです。1人につき1回のみ可能です。\" }",
    4, MONO_F, 90)

# ════════════════════════════════════════════════════════════════════
# Sheet 10: 8.認証認可
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("8.認証認可", "EF4444")
widths(ws, 24, 28, 52)

r = 2
r = hdr(ws, r, "8. 認証・認可設計", 3)
r = gap(r)
r = sub(ws, r, "8.1 認証フロー（モック JPKI）", 3)
r = para(ws, r,
    "(1) ユーザ /login にアクセス\n"
    "(2) サーバはペルソナリストを匿名ラベルのみで返す\n"
    "(3) ユーザがペルソナを選択 + カード読み取り → 1.8 秒のアニメ\n"
    "(4) 緑ボックスで display_name を初出\n"
    "(5) 「ログインしてダッシュボードへ」クリック → Server Action 起動\n"
    "(6) signIn(\"jpki-mock\", {personaKey}) → authorize() → hashCertificateSerial\n"
    "(7) prisma.user.upsert → JWT セッショントークン\n"
    "(8) リダイレクト /",
    3, MONO_F, 170)
r = gap(r)

r = sub(ws, r, "8.2 認証フロー（xID 実機・スタブ）", 3)
r = para(ws, r,
    "(1) ユーザが「実機で認証する」をクリック\n"
    "(2) GET /api/auth/signin/xid → NextAuth が認可リクエスト (PKCE)\n"
    "(3) xID 側で PIN 入力 + 証明書読み取り\n"
    "(4) コードコールバック → トークン交換 → id_token 受領\n"
    "(5) profile(): id = sha256(\"XID:\" + sub)、resident_code = municipality_code\n"
    "(6) signIn コールバック内で prisma.user.upsert\n"
    "(7) JWT セッション + リダイレクト /\n"
    "\n"
    "※ 環境変数 XID_ISSUER / XID_CLIENT_ID / XID_CLIENT_SECRET の 3 つが揃った時のみ有効化。",
    3, MONO_F, 200)
r = gap(r)

r = sub(ws, r, "8.3 多層防御", 3)
r = table(ws, r,
    ["層", "担当", "実装"],
    [
        ["1. Edge Middleware",           "URL ベースの粗い遮断",        "proxy.ts + authConfig.authorized"],
        ["2. Server Component",          "ページ表示前の再検証",        "redirect(\"/admin/forbidden\")"],
        ["3. Server Action / Route H.",  "操作実行前の最終チェック",    "isAdmin() + 自治体一致"],
    ],
    row_height=28,
)
r = gap(r)

r = sub(ws, r, "8.4 ロール", 3)
r = table(ws, r,
    ["ロール", "識別方法", "権限"],
    [
        ["未認証",      "session === null",                          "公開ページのみ閲覧"],
        ["一般住民",    "session.user.id && ¬admin",                 "自治体内提案の表示・投稿・投票・コメント"],
        ["職員",        "session.user.id ∈ ADMIN_USER_IDS",          "自治体内提案の管理・出力・検証"],
    ],
    row_height=28,
)
r = gap(r)

r = sub(ws, r, "8.5 セッション戦略", 3)
r = para(ws, r,
    "JWT 戦略を採用（session: { strategy: \"jwt\" }）\n"
    "理由：\n"
    " ・ サーバ間の状態共有不要 → 将来の水平スケール時に有利\n"
    " ・ DB セッション読み出しレイテンシなし\n"
    "拡張クレーム： uid, resident_code, is_verified",
    3, BODY, 110)

# ════════════════════════════════════════════════════════════════════
# Sheet 11: 9.マルチテナント
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("9.マルチテナント", "06B6D4")
widths(ws, 30, 70)

r = 2
r = hdr(ws, r, "9. マルチテナント設計", 2)
r = gap(r)
r = sub(ws, r, "9.1 テナント識別", 2)
r = para(ws, r,
    "テナント = 自治体 (municipality)。ユーザの resident_code（6 桁）がテナント識別子となる。\n"
    " ・ 単一プラットフォーム上で複数自治体が共存\n"
    " ・ 自治体の追加は MUNICIPALITIES マスタへの 1 行追加 + 必要なら職員ペルソナ\n"
    " ・ DB は単一インスタンス（行レベルで分離）",
    2, BODY, 90)
r = gap(r)

r = sub(ws, r, "9.2 テナント分離の階層", 2)
r = table(ws, r,
    ["層", "分離手段"],
    [
        ["UI フィルタ",       "提案一覧で where: { resident_code: session.user.resident_code }"],
        ["投稿時",            "Server Action が resident_code を session から自動付与（クライアント注入不可）"],
        ["投票・署名時",      "proposal.resident_code === session.user.resident_code を必須化"],
        ["管理時",            "管理 Action 内で同上のチェック"],
        ["エクスポート時",    "Route Handler 内で同上"],
    ],
    row_height=32,
)
r = gap(r)

r = sub(ws, r, "9.3 越境シナリオの扱い", 2)
r = table(ws, r,
    ["シナリオ", "挙動"],
    [
        ["千代田区民が盛岡市の URL を直接叩く（閲覧）",          "公開情報として閲覧可（透明性重視）"],
        ["千代田区民が盛岡市の請願に署名しようとする",            "Server Action が 403 相当のエラー返却"],
        ["千代田区職員が盛岡市の管理画面を URL 直叩き",            "一覧から消える + 個別ページは notFound()"],
    ],
    row_height=28,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 12: 10.セキュリティ
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("10.セキュリティ", "DC2626")
widths(ws, 24, 76)

r = 2
r = hdr(ws, r, "10. セキュリティ設計", 2)
r = gap(r)
r = sub(ws, r, "10.1 個人情報の最小化", 2)
r = para(ws, r,
    "■ マイナンバーは取得しない\n"
    "JPKI から取得可能な情報のうち、本システムが実装上参照しないもの：\n"
    " ・ マイナンバー（12 桁）\n"
    " ・ 性別・生年月日（プライバシー上不要）\n"
    "\n"
    "■ 不可逆ハッシュ化\n"
    "証明書シリアル番号 → SHA-256 → User.id に保存（元値復元は計算量的に不可能）\n"
    "実装: src/lib/jpki-mock.ts hashCertificateSerial()\n"
    "\n"
    "■ 表示時のフィルタ\n"
    " ・ コメント一覧: 「ある住民」+ 認証済みバッジ + 相対日時のみ\n"
    " ・ 監査ログ:   操作者 ID 先頭 8 文字 + …\n"
    " ・ CSV:        user_id_hashed（既にハッシュ済）\n"
    " ・ PDF:        統計値のみ。個別の署名者情報は含まれない",
    2, BODY, 270)
r = gap(r)

r = sub(ws, r, "10.2 改ざん検知 (HMAC-SHA256)", 2)
r = table(ws, r,
    ["対象", "HMAC の計算入力"],
    [
        ["CSV",
         "ファイル末尾に # HMAC-SHA256: <64桁hex> 行を追加。\n"
         "入力: BOM + プレアンブル + ヘッダー + データ行（\\r\\n 終端）"],
        ["PDF",
         "Document の Keywords metadata に hmac=<hex>;proposalId=<id>;issuedAt=<ISO> を設定。\n"
         "フッターにも文書検証用ハッシュとして hex を印字。\n"
         "入力: buildHmacPayload() が返す正規化 JSON （キー昇順、ISO 化）"],
    ],
    row_height=90,
)
r = gap(r)

r = sub(ws, r, "10.3 暗号強度の根拠", 2)
r = para(ws, r,
    " ・ 鍵は EXPORT_HMAC_SECRET（環境変数、サーバのみ保持）\n"
    " ・ HMAC は衝突困難： 32 バイトキー + SHA-256 で実質的に偽造不可\n"
    " ・ 比較は crypto.timingSafeEqual でタイミング攻撃耐性",
    2, BODY, 80)
r = gap(r)

r = sub(ws, r, "10.4 一般攻撃への対策", 2)
r = table(ws, r,
    ["攻撃カテゴリ", "対策"],
    [
        ["SQL インジェクション", "Prisma を介すためすべてパラメータ化クエリ。文字列連結 SQL は存在しない"],
        ["CSRF",                "Server Actions は Next.js が Action ID で検証、OIDC は PKCE + state"],
        ["XSS",                 "React デフォルトエスケープ依拠、dangerouslySetInnerHTML 未使用"],
        ["タイミング攻撃",      "HMAC 比較は timingSafeEqual で固定時間"],
    ],
    row_height=28,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 13: 11.監査証跡
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("11.監査証跡", "DC2626")
widths(ws, 26, 36, 38)

r = 2
r = hdr(ws, r, "11. 監査・証跡設計", 3)
r = gap(r)
r = sub(ws, r, "11.1 監査ログ三層", 3)
r = table(ws, r,
    ["テーブル", "何を記録", "監査が答えられる質問"],
    [
        ["AuditLog",        "提案ステータス変更",  "誰が・いつ・どの状態に変えたか"],
        ["IssuanceRecord",  "CSV/PDF 発行",        "いつ・誰が・どんな内容で文書を出したか"],
        ["VerificationLog", "検証実行",            "誰が・いつ・どのファイルを・どう判定したか"],
    ],
    row_height=30,
)
r = gap(r)

r = sub(ws, r, "11.2 Append-only 原則", 3)
r = para(ws, r,
    " ・ いずれのテーブルもアプリ層から UPDATE / DELETE 関数を提供しない\n"
    " ・ 「動かない記録」が積み上がることで内部不正・運用ミスの追跡が可能になる\n"
    " ・ 本番化時は PostgreSQL の ROW-level security + トリガでの UPDATE 拒否を併用推奨",
    3, BODY, 80)
r = gap(r)

r = sub(ws, r, "11.3 トランザクションによる原子性", 3)
r = para(ws, r,
    "ステータス変更時の処理：\n"
    "await prisma.$transaction(async (tx) => {\n"
    "  const before = await tx.proposal.findUnique(...)\n"
    "  if (before.status === newStatus) return  // no-op\n"
    "  await tx.proposal.update(...)\n"
    "  await tx.auditLog.create({...})  // ← 必ずペアで発行\n"
    "})\n"
    "\n"
    "部分失敗で「変更したのに監査ログがない」状態は構造上発生しない。",
    3, MONO_F, 180)
r = gap(r)

r = sub(ws, r, "11.4 検証可能性の確保", 3)
r = para(ws, r,
    "CSV / PDF の HMAC を埋め込み + 発行台帳保持により、\n"
    "行政側受領部門が「この文書は本当に発行された原本か」を\n"
    "追加サーバなしで証明できる。\n"
    "（/admin/verify ページにアップロードするだけで判定）",
    3, BODY, 110)

# ════════════════════════════════════════════════════════════════════
# Sheet 14: 12.非機能要件
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("12.非機能要件", "0EA5E9")
widths(ws, 26, 26, 48)

r = 2
r = hdr(ws, r, "12. 非機能要件", 3)
r = gap(r)
r = sub(ws, r, "12.1 性能", 3)
r = table(ws, r,
    ["指標", "目標値（プロト）", "本番化時の目安"],
    [
        ["ページ初期表示",        "< 2 秒",                                 "< 1 秒"],
        ["Server Action 応答",    "< 500ms",                                "< 200ms"],
        ["PDF 生成",              "< 3 秒（初回フォントロード含む）",       "< 2 秒"],
        ["HMAC 計算",             "< 10ms / KB",                            "同左"],
    ],
    row_height=26,
)
r = gap(r)

r = sub(ws, r, "12.2 可用性", 3)
r = para(ws, r,
    "・ プロト: ローカル単一プロセス\n"
    "・ 本番想定: SLA 99.5%（業務系として標準）",
    3, BODY, 50)
r = gap(r)

r = sub(ws, r, "12.3 スケーラビリティ", 3)
r = table(ws, r,
    ["軸", "限界（現状）", "拡張方向"],
    [
        ["自治体数",     "制限なし（マスタ追加のみ）",  "全国 1,718 自治体まで対応可能設計"],
        ["同時接続",     "単一 Node プロセス",          "Vercel / Cloud Run + DB プール"],
        ["データ量",     "SQLite ファイル",              "PostgreSQL（100 万署名/年でも余裕）"],
    ],
    row_height=28,
)
r = gap(r)

r = sub(ws, r, "12.4 アクセシビリティ", 3)
r = para(ws, r,
    "・ semantic HTML（button, form, ol 等）\n"
    "・ aria-label を主要操作に付与\n"
    "・ カラーコントラスト：ダークモード対応\n"
    "・ aria-live で動的領域を通知",
    3, BODY, 90)

# ════════════════════════════════════════════════════════════════════
# Sheet 15: 13.デプロイ運用
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("13.デプロイ運用", "0EA5E9")
widths(ws, 26, 26, 48)

r = 2
r = hdr(ws, r, "13. デプロイ・運用設計", 3)
r = gap(r)
r = sub(ws, r, "13.1 環境構成", 3)
r = table(ws, r,
    ["環境", "用途 / DB", "認証"],
    [
        ["ローカル",        "開発 / SQLite",      "JPKI モック"],
        ["ステージング",    "検証 / PostgreSQL",  "xID テストエンドポイント"],
        ["本番",            "運用 / PostgreSQL",  "xID 本番エンドポイント"],
    ],
    row_height=24,
)
r = gap(r)

r = sub(ws, r, "13.2 環境変数", 3)
r = table(ws, r,
    ["変数名", "用途", "機微度"],
    [
        ["DATABASE_URL",         "DB 接続文字列",             "高"],
        ["AUTH_SECRET",          "NextAuth JWT 署名鍵",        "最高"],
        ["AUTH_TRUST_HOST",      "リバプロ越し動作",           "中"],
        ["EXPORT_HMAC_SECRET",   "CSV/PDF 改ざん検知鍵",       "最高"],
        ["XID_ISSUER",           "xID OIDC discovery URL",     "中"],
        ["XID_CLIENT_ID",        "xID クライアント ID",        "高"],
        ["XID_CLIENT_SECRET",    "xID クライアント秘密",       "最高"],
    ],
    row_height=24,
)
r = gap(r)

r = sub(ws, r, "13.3 本番デプロイ手順", 3)
r = para(ws, r,
    "1. PostgreSQL を準備し DATABASE_URL を発行\n"
    "2. AUTH_SECRET / EXPORT_HMAC_SECRET を openssl rand -base64 32 で生成、KMS に保管\n"
    "3. xID クライアント登録・コールバック URL を https://<domain>/api/auth/callback/xid で登録\n"
    "4. .env を本番値で構成\n"
    "5. npx prisma migrate deploy でスキーマ反映（dev ではなく deploy を使用）\n"
    "6. npm run build → デプロイ（Vercel / Cloud Run / 自社環境）\n"
    "7. 動作確認チェックリスト:\n"
    "   □ xID で実機ログイン成功\n"
    "   □ 提案投稿 → 署名 → CSV/PDF 発行\n"
    "   □ 検証ページで OK 判定\n"
    "   □ 監査ログ・発行台帳・検証履歴に行が積まれている",
    3, BODY, 230)
r = gap(r)

r = sub(ws, r, "13.4 バックアップ", 3)
r = para(ws, r,
    "・ DB: 日次フルダンプ + WAL 連続バックアップ\n"
    "・ IssuanceRecord / AuditLog / VerificationLog: 別ストレージへ二重保管推奨（証跡性強化）",
    3, BODY, 60)
r = gap(r)

r = sub(ws, r, "13.5 ログ・モニタリング", 3)
r = para(ws, r,
    "・ アプリケーションログ: 構造化 JSON（請願 ID、操作者ハッシュ先頭、結果）\n"
    "・ 個人情報は出力しない\n"
    "・ アラート: 5xx 急増・PDF 生成失敗・HMAC 検証失敗多発",
    3, BODY, 80)

# ════════════════════════════════════════════════════════════════════
# Sheet 16: 14.制約とRM
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("14.制約とRM", "78716C")
widths(ws, 26, 34, 40)

r = 2
r = hdr(ws, r, "14. 既知の制約と本番化ロードマップ", 3)
r = gap(r)
r = sub(ws, r, "14.1 既知の制約", 3)
r = table(ws, r,
    ["制約", "影響", "対策"],
    [
        ["SQLite はマルチプロセス書込不向き",        "本番スループット限界",      "PostgreSQL（接続文字列のみ変更可）"],
        ["JPKI 認証はモック",                          "実機接続不可",              "xID 接続スタブを §8.2 で用意済み"],
        ["HMAC は共通鍵",                              "受領者にも鍵共有が必要",    "公開鍵署名（Ed25519）への移行を検討"],
        ["監査ログはアプリ層 append-only",             "DB 直叩きで改ざん可",       "WORM ストレージへの二重書きで強化"],
        ["コメントの通報・モデレーションが未実装",   "炎上時の運用負荷",          "フェーズ 2 機能として追加可能"],
    ],
    row_height=32,
)
r = gap(r)

r = sub(ws, r, "14.2 本番化ロードマップ", 3)
r = table(ws, r,
    ["フェーズ", "期間目安", "主要施策"],
    [
        ["フェーズ 1: 接続化",    "1〜2 か月", "xID 連携、PostgreSQL 移行、Vercel / Cloud Run デプロイ"],
        ["フェーズ 2: 統制強化",  "1〜2 か月", "KMS 統合、WORM 監査ログ、RBAC テーブル化、MFA 必須化"],
        ["フェーズ 3: 暗号強化",  "1 か月",    "HMAC → Ed25519 公開鍵署名へ移行（受領側が秘密を持たず検証可）"],
        ["フェーズ 4: 運用機能",  "2 か月",    "コメント通報、職員モデレーション UI、メール通知"],
        ["フェーズ 5: 横展開",    "継続",      "多言語化、モバイルアプリ、地理ヒートマップ"],
    ],
    row_height=32,
)

# ════════════════════════════════════════════════════════════════════
# Sheet 17: 付録
# ════════════════════════════════════════════════════════════════════
ws = new_sheet("付録", "64748B")
widths(ws, 26, 70)

r = 2
r = hdr(ws, r, "付録 A. 用語集", 2)
r = gap(r)
r = table(ws, r,
    ["用語", "説明"],
    [
        ["JPKI",            "Japanese Public Key Infrastructure。公的個人認証サービス"],
        ["マイナンバー",     "個人番号（12 桁）。本システムでは取扱なし"],
        ["署名用電子証明書", "マイナンバーカードに格納される PKI 証明書"],
        ["HMAC",            "Hash-based Message Authentication Code。共通鍵による完全性検証"],
        ["OIDC",            "OpenID Connect。OAuth 2.0 上のアイデンティティレイヤ"],
        ["xID",             "OIDC 経由で JPKI 認証を提供するサービスの一例"],
        ["append-only",     "追記のみ可能、変更・削除不可のテーブル運用方針"],
        ["canonical JSON",  "キー順序を正規化した JSON。ハッシュ入力の決定性確保"],
        ["PKCE",            "Proof Key for Code Exchange。OAuth/OIDC のセキュリティ拡張"],
        ["WORM",            "Write Once Read Many。改ざん不可ストレージ"],
        ["RBAC",            "Role-Based Access Control。役割ベースの権限制御"],
        ["JIS X 0402",      "全国地方公共団体コードの JIS 規格（6 桁）"],
    ],
    row_height=24,
)
r = gap(r, 2)

r = hdr(ws, r, "付録 B. 自治体マスタ（現プロト）", 2)
r = gap(r)
r = table(ws, r,
    ["コード", "表示名 / 備考"],
    [
        ["011002", "北海道札幌市"],
        ["131016", "東京都千代田区 ★主要デモ自治体"],
        ["131024", "東京都中央区"],
        ["131032", "東京都港区"],
        ["132012", "岩手県盛岡市（デモ用コード。実コード 032018）"],
        ["271004", "大阪府大阪市"],
        ["401307", "福岡県福岡市"],
    ],
    row_height=22,
)
r = gap(r)
r = para(ws, r,
    "正式運用時は JIS X 0402 全件のマスタを別テーブル化し、地方公共団体情報システム機構（J-LIS）が公表する最新情報を参照する想定。",
    2, NOTE, 40)
r = gap(r, 2)

r = hdr(ws, r, "付録 C. 改訂履歴", 2)
r = gap(r)
r = table(ws, r,
    ["版", "日付 / 改訂内容"],
    [
        ["1.0",   "2026-05-10 / 初版作成"],
    ],
    row_height=22,
)

# ════════════════════════════════════════════════════════════════════
# Save
# ════════════════════════════════════════════════════════════════════
import os
out = os.path.join(os.path.dirname(__file__), "設計書.xlsx")
wb.save(out)
print(f"WROTE: {out}")
print(f"SHEETS: {wb.sheetnames}")
